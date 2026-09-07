from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

base = Path(__file__).resolve().parent
src = base / 'campus-reserve-portfolio.txt'
out = base / 'output' / 'campus-reserve-portfolio.docx'
out.parent.mkdir(parents=True, exist_ok=True)

doc = Document()

styles = doc.styles
styles['Normal'].font.name = 'Calibri'
styles['Normal'].font.size = Pt(11)

# Title block
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Campus Reserve - Smart Facility Reservation System')
run.bold = True
run.font.size = Pt(16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Portfolio Builder 1')
run.italic = True
run.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Building Your System Administration Plan')
run.font.size = Pt(12)

lines = src.read_text(encoding='utf-8').splitlines()

i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue

    if line.startswith('Part A') or line.startswith('Part B') or line.startswith('Part C') or line.startswith('Part D') or line.startswith('Part E'):
        p = doc.add_paragraph()
        p.add_run(line.split('—', 1)[0].strip()).bold = True
        if '—' in line:
            p.add_run(' — ' + line.split('—', 1)[1].strip())
        i += 1

        table_rows = []
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                i += 1
                if table_rows:
                    break
                continue
            if nxt.startswith('Part '):
                break
            if nxt.startswith('|'):
                table_rows.append(nxt)
                i += 1
                continue
            # normal paragraph content
            if nxt.startswith('1.') or nxt.startswith('2.') or nxt.startswith('3.') or nxt.startswith('4.') or nxt.startswith('5.') or nxt.startswith('6.'):
                p = doc.add_paragraph()
                p.add_run(nxt).bold = True
            else:
                p = doc.add_paragraph()
                p.add_run(nxt)
            i += 1

        if table_rows:
            table = doc.add_table(rows=1, cols=2)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            headers = ['SDLC Phase', 'Administrative Activity']
            for c, h in enumerate(headers):
                table.cell(0, c).text = h
            for row_text in table_rows[2:]:
                cells = [c.strip() for c in row_text.strip('|').split('|')]
                if len(cells) >= 2:
                    row_cells = table.add_row().cells
                    row_cells[0].text = cells[0]
                    row_cells[1].text = cells[1]
            continue

        continue

    # fallback paragraph
    if line:
        p = doc.add_paragraph()
        p.add_run(line)
    i += 1

# Save document
doc.save(out)
print(f'DOCX created: {out}')
